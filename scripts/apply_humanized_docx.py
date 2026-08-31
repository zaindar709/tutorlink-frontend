"""
Merge humanized paragraph batches into a copy of the TutorLink FYP docx.
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document

ROOT = Path(r"E:\FYP\tutorlink-frontend")
SRC = Path(r"C:\Users\zaind\Downloads\TUTORLINK Final Documentation (1).docx")
OUT = Path(r"C:\Users\zaind\Downloads\TUTORLINK_Final_Documentation_HUMANIZED.docx")
OUT_ALT = ROOT / "TUTORLINK_Final_Documentation_HUMANIZED.docx"

BATCHES = [
    ROOT / "tmp_human_v2_batch1.json",
    ROOT / "tmp_human_v2_batch2.json",
    ROOT / "tmp_human_v2_batch3.json",
    ROOT / "tmp_human_v2_batch4.json",
    ROOT / "tmp_human_v2_medium.json",
]


def main() -> None:
    mapping: dict[int, str] = {}
    for path in BATCHES:
        if not path.exists():
            print("missing", path.name)
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        for item in data:
            mapping[int(item["i"])] = item["text"]
        print("loaded", path.name, len(data))

    print("total mapped", len(mapping))
    if not mapping:
        raise SystemExit("No humanized batches found")

    shutil.copy2(SRC, OUT)
    doc = Document(str(OUT))
    changed = 0
    for idx, new_text in mapping.items():
        if idx < 0 or idx >= len(doc.paragraphs):
            print("skip bad index", idx)
            continue
        para = doc.paragraphs[idx]
        if para.text == new_text:
            continue
        # Preserve runs formatting where possible: replace full paragraph text
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = new_text
        changed += 1

    doc.save(str(OUT))
    shutil.copy2(OUT, OUT_ALT)
    print("saved", OUT)
    print("copy", OUT_ALT)
    print("paragraphs changed", changed)


if __name__ == "__main__":
    main()
